"""The committed outputs in results/ are consistent with the code: tables regenerate byte for byte, and a sample of
the exact-model results recomputes. (A full regeneration takes about two minutes: python -m jpdm results.)"""
import json
from pathlib import Path

import pytest

from jpdm import exit_runs, tables

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"


def test_tables_regenerate_identically(tmp_path):
    tables.main(RESULTS, tmp_path)
    for f in (RESULTS / "tables").glob("*.md"):
        assert (tmp_path / f.name).read_text() == f.read_text(), f.name
    assert json.loads((tmp_path / "numbers.json").read_text()) == json.loads((RESULTS / "tables" / "numbers.json").read_text())


@pytest.mark.parametrize("key", ["dairy_2026", "eggs_2019-21"])
def test_full_loss_sample_recomputes(key):
    saved = json.loads((RESULTS / "full_loss.json").read_text())[key]
    br = next(b for b in exit_runs.BRANCHES if b.key == key)
    from jpdm import exact_model as X
    ex = X.exact_for(br)
    for stream, entry in saved.items():
        r = ex.solve(cap={stream: entry["cap"]}, eps_s=br.eps_s_lr)
        assert abs((r["A"] - 1.0) - entry["LR"]["A"]) < 1e-9, stream


def test_workbook_builds(tmp_path):
    from openpyxl import load_workbook
    from jpdm import workbook
    path = tmp_path / "model.xlsx"
    workbook.build(path, RESULTS)
    wb = load_workbook(path)
    assert wb.sheetnames == ["README", "Summary", "Assumptions", "Dairy", "Beef", "Eggs", "Broiler", "Pork", "Pollock",
                             "Salmon", "Sensitivity", "Ranges", "Exit", "Sources"]
    formulas = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas > 2000


def test_docs_regenerate_identically():
    from jpdm import docs
    assert docs.parameters_md() == (ROOT / "docs" / "parameters.md").read_text()
    assert docs.sources_md() == (ROOT / "docs" / "sources.md").read_text()
